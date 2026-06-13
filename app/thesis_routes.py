from __future__ import annotations

import re
import json
from io import BytesIO

from fastapi import APIRouter, HTTPException, Request, Response
from fastapi.responses import HTMLResponse, PlainTextResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel, Field, field_validator

from app.config import THESIS_MAX_EVIDENCE, THESIS_MAX_PER_WORKSPACE, THESIS_MODEL
from app.database import (
    add_article_evidence,
    add_snapshot_evidence,
    count_thesis_evidence,
    count_theses,
    create_thesis,
    delete_thesis_evidence,
    get_thesis,
    list_theses,
    list_thesis_evidence,
    save_thesis_memo,
    save_thesis_section_data,
    save_thesis_sections,
    search_articles_for_evidence,
    update_thesis,
)
from app.theses import (
    DRAFT_SECTION_KEYS,
    INSUFFICIENT_EVIDENCE_MESSAGES,
    clean_memo_display,
    generate_investment_memo,
    generate_thesis_sections,
    normalize_sections,
    present_evidence,
    regenerate_thesis_section,
    section_research_query,
    thesis_evidence_coverage,
    thesis_research_seed,
    upgrade_legacy_sections,
)
from app.thesis_research import (
    add_candidate_to_thesis,
    research_thesis_evidence,
    resolve_candidate,
)
from app.workspace import set_workspace_cookie, workspace_for_request


router = APIRouter()
templates = Jinja2Templates(directory="app/templates")


class ThesisCreate(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    core_claim: str = Field(min_length=1, max_length=3000)


class ThesisUpdate(ThesisCreate):
    domain: str | None = Field(default=None, max_length=200)
    research_question: str | None = Field(default=None, max_length=3000)
    initial_view: str | None = Field(default=None, max_length=3000)


class ThesisDraftCreate(BaseModel):
    domain: str = Field(min_length=1, max_length=200)
    research_question: str | None = Field(default=None, max_length=3000)
    initial_view: str | None = Field(default=None, max_length=3000)


class ThesisSectionsUpdate(BaseModel):
    sections: dict


class EvidenceResearchRequest(BaseModel):
    query: str = Field(min_length=2, max_length=1500)


class EvidenceCandidateAction(BaseModel):
    ref: str = Field(min_length=3, max_length=6000)
    classification: str = "CONTEXT"
    relationship_explanation: str = Field(default="", max_length=2000)
    strength: str = "medium"
    source_credibility: str = "medium"
    ai_recommendation: str = "consider"
    evidence_state: str = "added"


class ArticleEvidenceRequest(BaseModel):
    article_ids: list[int] = Field(min_length=1, max_length=5)


class SnapshotEvidenceRequest(BaseModel):
    evidence_type: str
    title: str = Field(min_length=1, max_length=250)
    excerpt: str = Field(min_length=1, max_length=12000)
    source: str | None = Field(default=None, max_length=200)
    url: str | None = Field(default=None, max_length=2000)
    note: str | None = Field(default=None, max_length=2000)
    article_ids: list[int] = Field(default_factory=list, max_length=10)
    metadata: dict = Field(default_factory=dict)

    @field_validator("url")
    @classmethod
    def validate_url(cls, value: str | None) -> str | None:
        if value and not value.startswith(("https://", "http://")):
            raise ValueError("URL must use http or https.")
        return value


def owned_thesis_or_404(thesis_id: int, workspace_id: str) -> dict:
    thesis = get_thesis(thesis_id, workspace_id)
    if not thesis:
        raise HTTPException(status_code=404, detail="Thesis not found.")
    return thesis


def add_researched_evidence(
    thesis: dict, workspace_id: str, query: str, *, limit: int
) -> list[dict]:
    existing = list_thesis_evidence(thesis["id"], workspace_id)
    existing_urls = {item.get("url") for item in existing if item.get("url")}
    added = []
    candidates = [
        item for item in research_thesis_evidence(query, workspace_id)
        if item["ai_recommendation"] != "low_priority"
    ]
    web_candidates = [item for item in candidates if item["evidence_type"] == "web"]
    other_candidates = [item for item in candidates if item["evidence_type"] != "web"]
    web_quota = min(len(web_candidates), max(1, limit // 3))
    ordered = web_candidates[:web_quota] + other_candidates + web_candidates[web_quota:]
    for candidate_data in ordered:
        if len(added) >= limit or count_thesis_evidence(thesis["id"], workspace_id) >= THESIS_MAX_EVIDENCE:
            break
        candidate = resolve_candidate(candidate_data["ref"], workspace_id)
        if not candidate or (candidate.url and candidate.url in existing_urls):
            continue
        item = add_candidate_to_thesis(
            thesis["id"],
            workspace_id,
            candidate,
            classification=candidate_data["classification"],
            relationship_explanation=candidate_data["relationship_explanation"],
            strength=candidate_data["strength"],
            source_credibility=candidate_data["source_credibility"],
            ai_recommendation=candidate_data["ai_recommendation"],
            evidence_state="added",
        )
        if item:
            added.append(item)
            if candidate.url:
                existing_urls.add(candidate.url)
    return added


def build_initial_evidence_library(
    thesis: dict, workspace_id: str, *, limit: int = 30
) -> list[dict]:
    added = add_researched_evidence(
        thesis, workspace_id, thesis_research_seed(thesis), limit=min(12, limit)
    )
    evidence = list_thesis_evidence(thesis["id"], workspace_id)
    coverage = thesis_evidence_coverage(thesis, evidence)
    priorities = (
        "core_claim", "why_now", "market_structure", "counter_evidence",
        "key_risks", "potential_winners", "market_drivers",
    )
    for key in priorities:
        if len(added) >= limit:
            break
        if coverage[key]["status"] != "insufficient":
            continue
        new_items = add_researched_evidence(
            thesis,
            workspace_id,
            section_research_query(thesis, key),
            limit=min(3, limit - len(added)),
        )
        added.extend(new_items)
        if new_items:
            evidence = list_thesis_evidence(thesis["id"], workspace_id)
            coverage = thesis_evidence_coverage(thesis, evidence)
    return added


def attach_cookie(
    response: Response, request: Request, workspace_id: str, is_new: bool
) -> None:
    if is_new:
        set_workspace_cookie(response, workspace_id, request)


@router.get("/theses", response_class=HTMLResponse)
def theses_page(request: Request) -> HTMLResponse:
    workspace_id, is_new = workspace_for_request(request)
    response = templates.TemplateResponse(
        request, "theses.html", {"theses": list_theses(workspace_id)}
    )
    attach_cookie(response, request, workspace_id, is_new)
    return response


@router.get("/thesis-builder", response_class=HTMLResponse)
def thesis_builder_page(request: Request) -> HTMLResponse:
    return theses_page(request)


@router.get("/theses/new", response_class=HTMLResponse)
def thesis_new_page(request: Request) -> HTMLResponse:
    workspace_id, is_new = workspace_for_request(request)
    response = templates.TemplateResponse(request, "thesis_new.html", {})
    attach_cookie(response, request, workspace_id, is_new)
    return response


@router.get("/theses/{thesis_id}", response_class=HTMLResponse)
def thesis_detail_page(request: Request, thesis_id: int) -> HTMLResponse:
    workspace_id, is_new = workspace_for_request(request)
    thesis = get_thesis(thesis_id, workspace_id)
    evidence = list_thesis_evidence(thesis_id, workspace_id) if thesis else []
    presented_evidence = [present_evidence(item) for item in evidence]
    sections = upgrade_legacy_sections(thesis.get("generated_sections")) if thesis else {}
    coverage = thesis_evidence_coverage(thesis, evidence) if thesis else {}
    response = templates.TemplateResponse(
        request,
        "thesis_detail.html",
        {
            "thesis": thesis,
            "evidence": presented_evidence,
            "sections": sections,
            "coverage": coverage,
            "memo_display": clean_memo_display(thesis.get("investment_memo")) if thesis else "",
            "section_keys": DRAFT_SECTION_KEYS,
            "max_evidence": THESIS_MAX_EVIDENCE,
        },
        status_code=200 if thesis else 404,
    )
    attach_cookie(response, request, workspace_id, is_new)
    return response


@router.get("/thesis-builder/{thesis_id}", response_class=HTMLResponse)
def thesis_builder_detail_page(request: Request, thesis_id: int) -> HTMLResponse:
    return thesis_detail_page(request, thesis_id)


@router.get("/api/theses")
def api_list_theses(request: Request, response: Response) -> list[dict]:
    workspace_id, is_new = workspace_for_request(request)
    attach_cookie(response, request, workspace_id, is_new)
    return list_theses(workspace_id)


@router.post("/api/theses", status_code=201)
def api_create_thesis(
    request: Request, response: Response, payload: ThesisCreate
) -> dict:
    workspace_id, is_new = workspace_for_request(request)
    title = payload.title.strip()
    core_claim = payload.core_claim.strip()
    if not title or not core_claim:
        raise HTTPException(status_code=400, detail="Title and core claim are required.")
    if count_theses(workspace_id) >= THESIS_MAX_PER_WORKSPACE:
        raise HTTPException(status_code=400, detail="Workspace thesis limit reached.")
    thesis = create_thesis(workspace_id, title, core_claim)
    attach_cookie(response, request, workspace_id, is_new)
    return thesis


@router.patch("/api/theses/{thesis_id}")
def api_update_thesis(
    request: Request, response: Response, thesis_id: int, payload: ThesisUpdate
) -> dict:
    workspace_id, is_new = workspace_for_request(request)
    existing = owned_thesis_or_404(thesis_id, workspace_id)
    title = payload.title.strip()
    core_claim = payload.core_claim.strip()
    if core_claim in INSUFFICIENT_EVIDENCE_MESSAGES:
        core_claim = (
            existing.get("initial_view")
            or (
                existing.get("core_claim")
                if existing.get("core_claim") not in INSUFFICIENT_EVIDENCE_MESSAGES
                else "待验证的投资判断"
            )
        )
    if not title or not core_claim:
        raise HTTPException(status_code=400, detail="Title and core claim are required.")
    thesis = update_thesis(
        thesis_id,
        workspace_id,
        title=title,
        core_claim=core_claim,
        domain=(
            payload.domain
            if "domain" in payload.model_fields_set
            else existing.get("domain")
        ),
        research_question=(
            payload.research_question
            if "research_question" in payload.model_fields_set
            else existing.get("research_question")
        ),
        initial_view=(
            payload.initial_view
            if "initial_view" in payload.model_fields_set
            else existing.get("initial_view")
        ),
    )
    if not thesis:
        raise HTTPException(status_code=404, detail="Thesis not found.")
    attach_cookie(response, request, workspace_id, is_new)
    return thesis


@router.post("/api/theses/draft", status_code=201)
def api_create_thesis_draft(
    request: Request, response: Response, payload: ThesisDraftCreate
) -> dict:
    workspace_id, is_new = workspace_for_request(request)
    if count_theses(workspace_id) >= THESIS_MAX_PER_WORKSPACE:
        raise HTTPException(status_code=400, detail="Workspace thesis limit reached.")
    domain = payload.domain.strip()
    question = (payload.research_question or "").strip()
    initial_view = (payload.initial_view or "").strip()
    seed_claim = initial_view or question or f"研究 {domain} 的结构性投资机会。"
    thesis = create_thesis(
        workspace_id,
        f"{domain} 投资论点",
        seed_claim,
        domain=domain,
        research_question=question or None,
        initial_view=initial_view or None,
    )
    attach_cookie(response, request, workspace_id, is_new)
    try:
        added = build_initial_evidence_library(thesis, workspace_id, limit=30)
        evidence = list_thesis_evidence(thesis["id"], workspace_id)
        sections = generate_thesis_sections(thesis, evidence)
        thesis = save_thesis_section_data(
            thesis["id"], workspace_id, sections, THESIS_MODEL
        ) or thesis
    except Exception as exc:
        return {"thesis": thesis, "sections": {}, "warning": str(exc)}
    return {
        "thesis": thesis,
        "sections": sections,
        "coverage": thesis_evidence_coverage(thesis, evidence),
        "added_evidence": added,
        "warning": None,
    }


@router.get("/api/theses/{thesis_id}/evidence/search")
def api_search_evidence(request: Request, thesis_id: int, q: str = "") -> list[dict]:
    workspace_id, _ = workspace_for_request(request)
    owned_thesis_or_404(thesis_id, workspace_id)
    query = q.strip()
    return search_articles_for_evidence(query[:200]) if len(query) >= 2 else []


@router.post("/api/theses/{thesis_id}/evidence/articles")
def api_add_article_evidence(
    request: Request, thesis_id: int, payload: ArticleEvidenceRequest
) -> dict:
    workspace_id, _ = workspace_for_request(request)
    owned_thesis_or_404(thesis_id, workspace_id)
    added = []
    for article_id in dict.fromkeys(payload.article_ids):
        if count_thesis_evidence(thesis_id, workspace_id) >= THESIS_MAX_EVIDENCE:
            break
        item = add_article_evidence(thesis_id, workspace_id, article_id)
        if item:
            added.append(item)
    if not added:
        raise HTTPException(status_code=400, detail="No evidence added or evidence limit reached.")
    return {
        "added": [present_evidence(item) for item in added],
        "evidence_count": count_thesis_evidence(thesis_id, workspace_id),
    }


@router.post("/api/theses/{thesis_id}/evidence/import")
def api_import_evidence(
    request: Request, thesis_id: int, payload: SnapshotEvidenceRequest
) -> dict:
    workspace_id, _ = workspace_for_request(request)
    owned_thesis_or_404(thesis_id, workspace_id)
    if payload.evidence_type not in {
        "chat",
        "comparison",
        "weekly_signal",
        "company_research",
        "manual",
    }:
        raise HTTPException(status_code=400, detail="Unsupported evidence type.")
    if len(json.dumps(payload.metadata)) > 5000:
        raise HTTPException(status_code=400, detail="Evidence metadata is too large.")
    if count_thesis_evidence(thesis_id, workspace_id) >= THESIS_MAX_EVIDENCE:
        raise HTTPException(status_code=400, detail="Evidence limit reached.")
    snapshot = add_snapshot_evidence(
        thesis_id,
        workspace_id,
        evidence_type=payload.evidence_type,
        title=payload.title.strip(),
        excerpt=payload.excerpt.strip(),
        source=payload.source,
        url=payload.url,
        note=payload.note,
        metadata=payload.metadata,
    )
    added_articles = []
    for article_id in dict.fromkeys(payload.article_ids):
        if count_thesis_evidence(thesis_id, workspace_id) >= THESIS_MAX_EVIDENCE:
            break
        item = add_article_evidence(thesis_id, workspace_id, article_id)
        if item:
            added_articles.append(item)
    return {"snapshot": snapshot, "articles": added_articles}


@router.post("/api/theses/{thesis_id}/evidence/research")
def api_research_evidence(
    request: Request, thesis_id: int, payload: EvidenceResearchRequest
) -> dict:
    workspace_id, _ = workspace_for_request(request)
    owned_thesis_or_404(thesis_id, workspace_id)
    try:
        candidates = research_thesis_evidence(payload.query.strip(), workspace_id)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    return {"query": payload.query.strip(), "evidence": candidates}


@router.post("/api/theses/{thesis_id}/evidence/candidates")
def api_add_evidence_candidate(
    request: Request, thesis_id: int, payload: EvidenceCandidateAction
) -> dict:
    workspace_id, _ = workspace_for_request(request)
    owned_thesis_or_404(thesis_id, workspace_id)
    if count_thesis_evidence(thesis_id, workspace_id) >= THESIS_MAX_EVIDENCE:
        raise HTTPException(status_code=400, detail="Evidence limit reached.")
    candidate = resolve_candidate(payload.ref, workspace_id)
    if not candidate:
        raise HTTPException(status_code=404, detail="Evidence candidate not found.")
    added = add_candidate_to_thesis(
        thesis_id,
        workspace_id,
        candidate,
        classification=payload.classification,
        relationship_explanation=payload.relationship_explanation,
        strength=payload.strength,
        source_credibility=payload.source_credibility,
        ai_recommendation=payload.ai_recommendation,
        evidence_state=payload.evidence_state,
    )
    if not added:
        raise HTTPException(status_code=400, detail="Unable to add evidence.")
    return {"evidence": present_evidence(added)}


@router.delete("/api/theses/{thesis_id}/evidence/{evidence_id}")
def api_delete_evidence(request: Request, thesis_id: int, evidence_id: int) -> dict:
    workspace_id, _ = workspace_for_request(request)
    if not delete_thesis_evidence(thesis_id, workspace_id, evidence_id):
        raise HTTPException(status_code=404, detail="Evidence not found.")
    return {"deleted": True}


@router.post("/api/theses/{thesis_id}/generate")
def api_generate_thesis(request: Request, thesis_id: int) -> dict:
    workspace_id, _ = workspace_for_request(request)
    thesis = owned_thesis_or_404(thesis_id, workspace_id)
    try:
        evidence = [
            item for item in list_thesis_evidence(thesis_id, workspace_id)
            if item.get("evidence_state") != "saved_for_later"
        ]
        sections = generate_thesis_sections(thesis, evidence)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    updated = save_thesis_sections(thesis_id, workspace_id, sections, THESIS_MODEL)
    return {
        "thesis": updated,
        "sections": sections,
        "coverage": thesis_evidence_coverage(thesis, evidence),
    }


@router.post("/api/theses/{thesis_id}/research/refresh")
def api_refresh_thesis_research(request: Request, thesis_id: int) -> dict:
    workspace_id, _ = workspace_for_request(request)
    thesis = owned_thesis_or_404(thesis_id, workspace_id)
    current_count = count_thesis_evidence(thesis_id, workspace_id)
    remaining = max(0, min(20, min(30, THESIS_MAX_EVIDENCE) - current_count))
    if not remaining:
        evidence = list_thesis_evidence(thesis_id, workspace_id)
        return {"added": [], "coverage": thesis_evidence_coverage(thesis, evidence)}
    try:
        added = build_initial_evidence_library(thesis, workspace_id, limit=remaining)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    evidence = list_thesis_evidence(thesis_id, workspace_id)
    return {
        "added": [present_evidence(item) for item in added],
        "coverage": thesis_evidence_coverage(thesis, evidence),
    }


@router.patch("/api/theses/{thesis_id}/sections")
def api_update_thesis_sections(
    request: Request, thesis_id: int, payload: ThesisSectionsUpdate
) -> dict:
    workspace_id, _ = workspace_for_request(request)
    owned_thesis_or_404(thesis_id, workspace_id)
    evidence = list_thesis_evidence(thesis_id, workspace_id)
    sections = normalize_sections(payload.sections, evidence)
    updated = save_thesis_section_data(thesis_id, workspace_id, sections)
    return {"thesis": updated, "sections": sections}


@router.post("/api/theses/{thesis_id}/sections/{section_key}/regenerate")
def api_regenerate_thesis_section(
    request: Request, thesis_id: int, section_key: str
) -> dict:
    workspace_id, _ = workspace_for_request(request)
    thesis = owned_thesis_or_404(thesis_id, workspace_id)
    if section_key not in DRAFT_SECTION_KEYS:
        raise HTTPException(status_code=404, detail="Unknown thesis section.")
    try:
        evidence = [
            item for item in list_thesis_evidence(thesis_id, workspace_id)
            if item.get("evidence_state") != "saved_for_later"
        ]
        section = regenerate_thesis_section(thesis, evidence, section_key)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    sections = upgrade_legacy_sections(thesis.get("generated_sections"))
    sections[section_key] = section
    updated = save_thesis_section_data(
        thesis_id, workspace_id, sections, THESIS_MODEL
    )
    return {
        "thesis": updated,
        "section_key": section_key,
        "section": section,
        "coverage": thesis_evidence_coverage(thesis, evidence)[section_key],
    }


@router.post("/api/theses/{thesis_id}/sections/{section_key}/evidence/research")
def api_research_section_evidence(
    request: Request, thesis_id: int, section_key: str
) -> dict:
    workspace_id, _ = workspace_for_request(request)
    thesis = owned_thesis_or_404(thesis_id, workspace_id)
    if section_key not in DRAFT_SECTION_KEYS:
        raise HTTPException(status_code=404, detail="Unknown thesis section.")
    query = section_research_query(thesis, section_key)
    try:
        candidates = research_thesis_evidence(query, workspace_id)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    return {"query": query, "section_key": section_key, "evidence": candidates}


@router.post("/api/theses/{thesis_id}/memo")
def api_generate_memo(request: Request, thesis_id: int) -> dict:
    workspace_id, _ = workspace_for_request(request)
    thesis = owned_thesis_or_404(thesis_id, workspace_id)
    evidence = list_thesis_evidence(thesis_id, workspace_id)
    try:
        memo = generate_investment_memo(thesis, evidence)
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=502,
            detail=f"Investment memo generation failed: {exc}",
        ) from exc
    updated = save_thesis_memo(thesis_id, workspace_id, memo, THESIS_MODEL)
    return {"thesis": updated, "memo": memo}


@router.get("/theses/{thesis_id}/export.md")
def export_thesis_memo(request: Request, thesis_id: int) -> PlainTextResponse:
    workspace_id, _ = workspace_for_request(request)
    thesis = owned_thesis_or_404(thesis_id, workspace_id)
    if not thesis.get("investment_memo"):
        raise HTTPException(status_code=400, detail="Generate the memo before export.")
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", thesis["title"]).strip("-").lower()
    return PlainTextResponse(
        thesis["investment_memo"],
        media_type="text/markdown",
        headers={
            "Content-Disposition": f'attachment; filename="{slug or "investment-thesis"}.md"'
        },
    )


@router.get("/theses/{thesis_id}/export.txt")
def export_thesis_memo_text(request: Request, thesis_id: int) -> PlainTextResponse:
    workspace_id, _ = workspace_for_request(request)
    thesis = owned_thesis_or_404(thesis_id, workspace_id)
    if not thesis.get("investment_memo"):
        raise HTTPException(status_code=400, detail="Generate the memo before export.")
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", thesis["title"]).strip("-").lower()
    return PlainTextResponse(
        thesis["investment_memo"],
        media_type="text/plain",
        headers={"Content-Disposition": f'attachment; filename="{slug or "investment-thesis"}.txt"'},
    )


@router.get("/theses/{thesis_id}/export.docx")
def export_thesis_memo_docx(request: Request, thesis_id: int) -> StreamingResponse:
    from docx import Document

    workspace_id, _ = workspace_for_request(request)
    thesis = owned_thesis_or_404(thesis_id, workspace_id)
    memo = thesis.get("investment_memo")
    if not memo:
        raise HTTPException(status_code=400, detail="Generate the memo before export.")
    document = Document()
    document.add_heading(thesis["title"], 0)
    for raw_line in memo.splitlines():
        line = raw_line.strip()
        if line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line:
            document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    output.seek(0)
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", thesis["title"]).strip("-").lower()
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{slug or "investment-thesis"}.docx"'},
    )
